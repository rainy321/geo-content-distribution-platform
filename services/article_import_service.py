from __future__ import annotations

import csv
import io
import re
from typing import Any, BinaryIO

from openpyxl import Workbook, load_workbook


MAX_IMPORT_ROWS = 200
MAX_IMPORT_BYTES = 2 * 1024 * 1024

_COLUMN_ALIASES = {
    "title": {"title", "标题", "文章标题"},
    "summary": {"summary", "摘要", "文章摘要"},
    "content": {"content", "正文", "文章正文", "内容"},
    "tags": {"tags", "标签", "关键词"},
    "status": {"status", "状态"},
}
_STATUS_ALIASES = {
    "": "draft",
    "draft": "draft",
    "草稿": "draft",
    "ready": "ready",
    "就绪": "ready",
}


class ArticleImportError(ValueError):
    def __init__(self, message: str, *, rows: list[dict[str, Any]] | None = None):
        super().__init__(message)
        self.rows = rows or []


def parse_article_import(
    stream: BinaryIO,
    *,
    filename: str,
) -> list[dict[str, Any]]:
    raw = stream.read(MAX_IMPORT_BYTES + 1)
    if len(raw) > MAX_IMPORT_BYTES:
        raise ArticleImportError("导入文件不能超过 2MB")
    if not raw:
        raise ArticleImportError("导入文件为空")

    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix == "xlsx":
        rows = _read_xlsx(raw)
    elif suffix == "csv":
        rows = _read_csv(raw)
    else:
        raise ArticleImportError("仅支持 .xlsx 或 .csv 文件")
    return _normalize_rows(rows)


def build_article_import_template() -> io.BytesIO:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "文章导入"
    worksheet.append(["标题", "摘要", "正文", "标签", "状态"])
    worksheet.append(
        [
            "企业如何选择 AI Agent",
            "说明选择企业智能体时需要关注的关键维度。",
            "## 先明确业务问题\n\n在选择方案前，应先确认需要解决的业务问题。",
            "AI Agent,企业智能体",
            "草稿",
        ]
    )
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = "A1:E2"
    worksheet.column_dimensions["A"].width = 30
    worksheet.column_dimensions["B"].width = 38
    worksheet.column_dimensions["C"].width = 72
    worksheet.column_dimensions["D"].width = 30
    worksheet.column_dimensions["E"].width = 12
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _read_xlsx(raw: bytes) -> list[list[Any]]:
    try:
        workbook = load_workbook(
            io.BytesIO(raw),
            read_only=True,
            data_only=True,
        )
        worksheet = workbook.active
        rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
        workbook.close()
        return rows
    except Exception as exc:
        raise ArticleImportError("Excel 文件无法读取或已损坏") from exc


def _read_csv(raw: bytes) -> list[list[str]]:
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ArticleImportError("CSV 必须使用 UTF-8 编码") from exc
    try:
        return list(csv.reader(io.StringIO(text)))
    except csv.Error as exc:
        raise ArticleImportError("CSV 文件格式无效") from exc


def _normalize_rows(rows: list[list[Any]]) -> list[dict[str, Any]]:
    if not rows:
        raise ArticleImportError("导入文件没有表头")
    headers = [_normalize_header(value) for value in rows[0]]
    column_indexes: dict[str, int] = {}
    for canonical, aliases in _COLUMN_ALIASES.items():
        for index, header in enumerate(headers):
            if header in aliases:
                column_indexes[canonical] = index
                break
    missing = [name for name in ("title", "content") if name not in column_indexes]
    if missing:
        labels = {"title": "标题", "content": "正文"}
        raise ArticleImportError(
            "导入表缺少必需列：" + "、".join(labels[item] for item in missing)
        )

    result: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    for row_number, row in enumerate(rows[1:], start=2):
        if not any(str(value or "").strip() for value in row):
            continue
        if len(result) + len(errors) >= MAX_IMPORT_ROWS:
            raise ArticleImportError(f"一次最多导入 {MAX_IMPORT_ROWS} 篇文章")
        values = {
            name: _cell(row, index)
            for name, index in column_indexes.items()
        }
        title = values.get("title", "")
        content = values.get("content", "")
        status_raw = values.get("status", "").lower()
        row_errors = []
        if not title:
            row_errors.append("标题为空")
        if len(title) > 100:
            row_errors.append("标题超过 100 字")
        if not content:
            row_errors.append("正文为空")
        if status_raw not in _STATUS_ALIASES:
            row_errors.append("状态仅支持草稿/draft/就绪/ready")
        if row_errors:
            errors.append({"row": row_number, "errors": row_errors})
            continue
        result.append(
            {
                "title": title,
                "summary": values.get("summary", "")[:500],
                "content": content,
                "tags": _split_tags(values.get("tags", "")),
                "status": _STATUS_ALIASES[status_raw],
            }
        )

    if errors:
        raise ArticleImportError("导入文件包含无效行", rows=errors)
    if not result:
        raise ArticleImportError("导入文件没有可创建的文章")
    return result


def _normalize_header(value: Any) -> str:
    return str(value or "").strip().lower().replace(" ", "")


def _cell(row: list[Any], index: int) -> str:
    if index >= len(row) or row[index] is None:
        return ""
    return str(row[index]).strip()


def _split_tags(value: str) -> list[str]:
    result = []
    for item in re.split(r"[,，;；\n]+", value):
        tag = item.strip().lstrip("#").strip()
        if tag and tag not in result:
            result.append(tag)
    return result[:20]
